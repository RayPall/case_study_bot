import streamlit as st
import requests
from docx import Document
import base64
import io
import re

st.set_page_config(page_title="Generátor případových studií", layout="centered")
st.title("📄 Generátor případových studií → JSON pro Make")

# 1) Výběr šablony
st.subheader("1. Vyber šablonu")
template = st.radio("Zvol verzi šablony:", ["Bright", "Soft"], horizontal=True)

# 2) Nahrání DOCX s obsahem
st.subheader("2. Nahraj .docx soubor s textem případové studie")
docx_file = st.file_uploader("Text případové studie", type="docx", key="docx")

# 3) Nahrání fotodokumentace (volitelné)
st.subheader("3. Nahraj fotodokumentaci (volitelné)")
images = st.file_uploader(
    "Obrázky (JPG, PNG)", type=["jpg", "jpeg", "png"], accept_multiple_files=True, key="images"
)

# 4) Webhook URL
st.subheader("4. Make webhook URL")
webhook_url = st.text_input("Vlož URL na Custom webhook v Make:", placeholder="https://hook.integromat.com/…")

# Pomocná funkce: DOCX → JSON placeholder struktura
def parse_docx_to_placeholders(docx_bytes):
    doc = Document(io.BytesIO(docx_bytes))
    data = {k: "" for k in [
        "predstaveni1", "predstaveni2", "vykrik1", "jmenoapozice", "citace1",
        "subheader",
        "cislo1", "prinos1headline", "prinos1",
        "cislo2", "prinos2headline", "prinos2",
        "cislo3", "prinos3headline", "prinos3",
        "popisreseni1", "popisreseni2",
        "nasereseni1", "nasereseni2",
        "vykrik2", "popisprinosu1", "popisprinosu2",
        "postimplementace1",
        "benefit1", "benefit2", "benefit3",
        "kontaktjmeno", "kontaktcislo", "contactmail"
    ]}

    current_section = None
    
    mapping = {
        "Zákazník": "predstaveni1",
        "Výzva": "predstaveni2",
        "Řešení": "nasereseni1",
        "Výsledky": "popisprinosu1",
        "O společnosti": "popisreseni1"
    }

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style and para.style.name and para.style.name.startswith("Heading"):
            current_section = mapping.get(text, None)
        elif current_section:
            if data[current_section] == "":
                data[current_section] = text
            else:
                data[current_section] += "\n" + text

    # Dummy naplnění dalších polí (reálně bys měl parsovat sofistikovaněji)
    if data["predstaveni1"] and not data["predstaveni2"]:
        parts = data["predstaveni1"].split(". ")
        if len(parts) > 1:
            data["predstaveni1"] = parts[0]
            data["predstaveni2"] = ". ".join(parts[1:])

    # Testovací vykriky, čísla a benefity – v praxi načti z konkrétních sekcí
    data["vykrik1"] = "Hlavní benefit projektu"
    data["vykrik2"] = "Výsledek: vyšší efektivita"
    data["cislo1"] = "1600 EUR"
    data["prinos1headline"] = "Měsíční úspora"
    data["prinos1"] = "Snížení nákladů o 1600 EUR měsíčně."
    data["cislo2"] = "60 %"
    data["prinos2headline"] = "Zrychlení načítání"
    data["prinos2"] = "Doba načítání dat se zkrátila o 60 %."
    data["cislo3"] = "800"
    data["prinos3headline"] = "Počet commitů"
    data["prinos3"] = "Během projektu bylo provedeno přes 800 commitů."
    data["popisreseni2"] = "Další část popisu přínosů."
    data["nasereseni2"] = "Druhá část implementace."
    data["popisprinosu2"] = "Další výsledky projektu."
    data["postimplementace1"] = "Plánuje se další rozvoj řešení."
    data["benefit1"] = "Významné provozní úspory"
    data["benefit2"] = "Zkrácení času načítání"
    data["benefit3"] = "Vyšší flexibilita datového provozu"
    data["kontaktjmeno"] = "John Doe"
    data["kontaktcislo"] = "+420 777 777 777"
    data["contactmail"] = "john.doe@example.com"

    return data

# 5) Zpracování → JSON
st.subheader("5. Zpracování a JSON náhled")
include_images = st.checkbox("Zahrnout obrázky do JSON (base64)")

template_selected = template.lower()

if st.button("🧠 Vytvořit JSON náhled"):
    if not docx_file:
        st.error("Nejprve prosím nahraj .docx soubor.")
    else:
        parsed_data = parse_docx_to_placeholders(docx_file.getvalue())
        payload = {
            "template": template_selected,
            "data": parsed_data,
        }
        if include_images and images:
            imgs = []
            for i, img in enumerate(images):
                b64 = base64.b64encode(img.getvalue()).decode("utf-8")
                imgs.append({
                    "filename": img.name,
                    "content_type": img.type,
                    "content_base64": b64,
                })
            payload["images"] = imgs
        st.success("✅ JSON připraven. Zkontroluj obsah níže.")
        st.json(payload)
        st.session_state["last_payload"] = payload

# 6) Odeslání na Make webhook (application/json)
if st.button("🚀 Odeslat JSON do Make"):
    payload = st.session_state.get("last_payload")
    if not payload:
        st.error("Nejprve vytvoř JSON náhled (tlačítko výše).")
    elif not webhook_url:
        st.error("Zadej prosím Make webhook URL.")
    else:
        try:
            resp = requests.post(webhook_url, json=payload, timeout=30)
            if 200 <= resp.status_code < 300:
                st.success(f"✅ Odesláno. HTTP {resp.status_code}")
                if resp.text:
                    st.code(resp.text)
            else:
                st.error(f"❌ Chyba při odesílání: HTTP {resp.status_code}")
                st.code(resp.text)
        except Exception as e:
            st.error(f"❌ Výjimka při odeslání: {e}")

st.caption("Pozn.: Aplikace generuje JSON podle sjednoceného seznamu placeholderů. V praxi doporučuji rozšířit parser, aby skutečně extrahoval všechny konkrétní části dokumentu.")
